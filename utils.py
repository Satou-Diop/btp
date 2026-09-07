from datetime import datetime
import database as db

def calculate_tache_avancement(qte_prevue, qte_realisee):
    """Calculer l'avancement d'une tâche"""
    if qte_prevue == 0:
        return 0
    return min(qte_realisee / qte_prevue, 1.0)

def calculate_uo_avancement(uo_id):
    """Calculer l'avancement d'une UO (moyenne des tâches)"""
    taches = db.get_taches_by_uo(uo_id)
    
    if not taches:
        return 0
    
    avancements = [
        calculate_tache_avancement(t['qte_prevue'], t['qte_realisee'])
        for t in taches
    ]
    
    return sum(avancements) / len(avancements) if avancements else 0

def get_uo_status(avancement):
    """Déterminer le statut d'une UO basé sur l'avancement"""
    if avancement == 0:
        return "Planifiée"
    elif avancement >= 1.0:
        return "Terminée"
    else:
        return "En cours"

def get_uo_details(uo_id):
    """Récupérer tous les détails d'une UO avec calculs"""
    uos = db.get_connection().cursor().execute(
        'SELECT * FROM unites_oeuvre WHERE id=?', (uo_id,)
    ).fetchone()
    
    if not uos:
        return None
    
    uo = dict(uos)
    avancement = calculate_uo_avancement(uo_id)
    
    uo['avancement'] = avancement
    uo['statut'] = get_uo_status(avancement)
    uo['ecart_budget'] = uo['budget'] - uo['depenses']
    uo['taux_depense'] = (uo['depenses'] / uo['budget'] * 100) if uo['budget'] > 0 else 0
    
    return uo

def get_chantier_details(chantier_id):
    """Récupérer tous les détails d'un chantier"""
    chantier = db.get_chantier(chantier_id)
    if not chantier:
        return None
    
    uos = db.get_uos_by_chantier(chantier_id)
    
    # Calculer les avancements et statuts
    for uo in uos:
        uo['avancement'] = calculate_uo_avancement(uo['id'])
        uo['statut'] = get_uo_status(uo['avancement'])
        uo['ecart_budget'] = uo['budget'] - uo['depenses']
        uo['taux_depense'] = (uo['depenses'] / uo['budget'] * 100) if uo['budget'] > 0 else 0
        
        # Ajouter les tâches
        uo['taches'] = db.get_taches_by_uo(uo['id'])
    
    chantier['uos'] = uos
    
    # Calculer les totaux du chantier
    chantier['nb_uos'] = len(uos)
    chantier['nb_taches'] = sum(len(uo['taches']) for uo in uos)
    chantier['nb_uos_terminees'] = sum(1 for uo in uos if uo['statut'] == 'Terminée')
    chantier['nb_uos_en_cours'] = sum(1 for uo in uos if uo['statut'] == 'En cours')
    
    # Avancement global = moyenne des UO
    if uos:
        chantier['avancement_global'] = sum(uo['avancement'] for uo in uos) / len(uos)
    else:
        chantier['avancement_global'] = 0
    
    # Budgets
    chantier['budget_total'] = sum(uo['budget'] for uo in uos)
    chantier['depenses_total'] = sum(uo['depenses'] for uo in uos)
    chantier['ecart_total'] = chantier['budget_total'] - chantier['depenses_total']
    
    # Durée
    if chantier['date_debut'] and chantier['date_fin']:
        try:
            start = datetime.fromisoformat(chantier['date_debut'])
            end = datetime.fromisoformat(chantier['date_fin'])
            chantier['duree_jours'] = (end - start).days
        except:
            chantier['duree_jours'] = 0
    else:
        chantier['duree_jours'] = 0
    
    return chantier

def get_all_chantiers_details():
    """Récupérer tous les chantiers avec leurs détails"""
    chantiers = db.get_all_chantiers()
    return [get_chantier_details(c['id']) for c in chantiers]

def export_chantier_to_excel(chantier_id, filename):
    """Exporter un chantier en Excel"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    
    chantier = get_chantier_details(chantier_id)
    if not chantier:
        return None
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chantier"
    
    # En-tête
    ws['A1'] = f"CHANTIER : {chantier['nom']}"
    ws['A1'].font = Font(bold=True, size=14)
    
    ws['A2'] = f"Lieu : {chantier['lieu']} | Chef : {chantier['chef_chantier']}"
    ws['A3'] = f"Avancement : {chantier['avancement_global']*100:.1f}%"
    
    # Tableau UO
    ws['A5'] = "Unité d'Œuvre"
    ws['B5'] = "Statut"
    ws['C5'] = "Avancement"
    ws['D5'] = "Budget"
    ws['E5'] = "Dépenses"
    ws['F5'] = "Écart"
    
    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws[f'{col}5'].font = Font(bold=True, color="FFFFFF")
        ws[f'{col}5'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    row = 6
    for uo in chantier['uos']:
        ws[f'A{row}'] = uo['nom']
        ws[f'B{row}'] = uo['statut']
        ws[f'C{row}'] = f"{uo['avancement']*100:.1f}%"
        ws[f'D{row}'] = uo['budget']
        ws[f'E{row}'] = uo['depenses']
        ws[f'F{row}'] = uo['ecart_budget']
        row += 1
    
    # Totaux
    ws[f'A{row+1}'] = "TOTAL"
    ws[f'A{row+1}'].font = Font(bold=True)
    ws[f'D{row+1}'] = chantier['budget_total']
    ws[f'E{row+1}'] = chantier['depenses_total']
    ws[f'F{row+1}'] = chantier['ecart_total']
    
    wb.save(filename)
    return filename
